import io
from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from models import PAYMENT_STATUS_LABELS

BRAND_FILL = PatternFill(start_color="1B2130", end_color="1B2130", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
TITLE_FONT = Font(bold=True, size=14, color="1B2130")
SUBTLE_FONT = Font(color="5B6472", size=10, italic=True)
SECTION_FONT = Font(bold=True, size=12, color="1B2130")
CURRENCY_FORMAT = '"₹"#,##0.00'


def _style_header_row(ws, row_idx, num_cols):
    for col in range(1, num_cols + 1):
        cell = ws.cell(row=row_idx, column=col)
        cell.fill = BRAND_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")


def _autosize(ws, headers, min_width=12, max_width=45):
    for idx, header in enumerate(headers, start=1):
        col_letter = get_column_letter(idx)
        ws.column_dimensions[col_letter].width = max(min_width, min(max_width, len(str(header)) + 4))


def build_expense_report(expenses, from_date=None, to_date=None):
    """Two sheets: Summary (aggregates, employee-wise breakdown) and
    Transactions (full detail row per expense). `expenses` is the same
    filtered list shown on the Admin Reports screen."""
    wb = Workbook()

    ws = wb.active
    ws.title = "Summary"

    ws["A1"] = "Final Mile Techies — Expense Report"
    ws["A1"].font = TITLE_FONT
    period = f"{from_date} to {to_date}" if from_date and to_date else "All dates"
    ws["A2"] = f"Period: {period}"
    ws["A2"].font = SUBTLE_FONT
    ws["A3"] = f"Generated: {datetime.now().strftime('%d %b %Y, %I:%M %p')}"
    ws["A3"].font = SUBTLE_FONT

    total_expenses = sum(float(e.amount) for e in expenses)
    logistics_total = sum(float(e.amount) for e in expenses if e.expense_type == "LOGISTICS")
    warehousing_total = sum(float(e.amount) for e in expenses if e.expense_type == "WAREHOUSING")

    row = 5
    ws.cell(row=row, column=1, value="Metric")
    ws.cell(row=row, column=2, value="Value")
    _style_header_row(ws, row, 2)
    row += 1

    for label, value, is_currency in [
        ("Total Expenses", total_expenses, True),
        ("Logistics Expenses", logistics_total, True),
        ("Warehousing Expenses", warehousing_total, True),
        ("Total Transactions", len(expenses), False),
    ]:
        ws.cell(row=row, column=1, value=label)
        cell = ws.cell(row=row, column=2, value=value)
        if is_currency:
            cell.number_format = CURRENCY_FORMAT
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="Employee-wise Expenses").font = SECTION_FONT
    row += 1
    ws.cell(row=row, column=1, value="Employee ID")
    ws.cell(row=row, column=2, value="Employee Name")
    ws.cell(row=row, column=3, value="Total Spent")
    _style_header_row(ws, row, 3)
    row += 1

    if expenses:
        df = pd.DataFrame(
            [
                {
                    "employee_id": e.user.employee_id,
                    "employee_name": e.user.name,
                    "amount": float(e.amount),
                }
                for e in expenses
            ]
        )
        grouped = (
            df.groupby(["employee_id", "employee_name"])["amount"]
            .sum()
            .reset_index()
            .sort_values("amount", ascending=False)
        )
        for _, r in grouped.iterrows():
            ws.cell(row=row, column=1, value=r["employee_id"])
            ws.cell(row=row, column=2, value=r["employee_name"])
            cell = ws.cell(row=row, column=3, value=float(r["amount"]))
            cell.number_format = CURRENCY_FORMAT
            row += 1

    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = 18

    ws2 = wb.create_sheet("Transactions")
    headers = [
        "Transaction ID", "Date", "Time", "Employee ID", "Employee Name",
        "Expense Type", "Docket Number", "Category", "Amount", "Purpose", "Reason",
        "Approved By", "Remarks", "UPI Reference Number", "Payment Status",
        "Rejection Reason", "Invoice File", "Payment Screenshot",
    ]
    for idx, header in enumerate(headers, start=1):
        ws2.cell(row=1, column=idx, value=header)
    _style_header_row(ws2, 1, len(headers))
    ws2.freeze_panes = "A2"

    for r_idx, e in enumerate(expenses, start=2):
        values = [
            e.transaction_id,
            e.created_at.strftime("%d-%m-%Y"),
            e.created_at.strftime("%I:%M %p"),
            e.user.employee_id,
            e.user.name,
            e.expense_type.title(),
            e.docket_no or "N/A",
            e.category or "N/A",
            float(e.amount),
            e.purpose,
            e.reason or "N/A",
            e.approved_by,
            e.remarks or "",
            e.upi_reference_no,
            PAYMENT_STATUS_LABELS.get(e.payment_status, e.payment_status),
            e.rejection_reason or "",
            e.invoice_file,
            e.payment_screenshot,
        ]
        for c_idx, val in enumerate(values, start=1):
            cell = ws2.cell(row=r_idx, column=c_idx, value=val)
            if c_idx == 9:
                cell.number_format = CURRENCY_FORMAT

    last_col = get_column_letter(len(headers))
    ws2.auto_filter.ref = f"A1:{last_col}{max(1, len(expenses) + 1)}"
    _autosize(ws2, headers)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def build_money_distribution_report(transactions):
    wb = Workbook()
    ws = wb.active
    ws.title = "Money Distribution"

    headers = [
        "Date", "Transaction ID", "Employee ID", "Employee Name",
        "Amount Added", "Purpose", "Remarks", "Created By",
    ]
    for idx, header in enumerate(headers, start=1):
        ws.cell(row=1, column=idx, value=header)
    _style_header_row(ws, 1, len(headers))
    ws.freeze_panes = "A2"

    for r_idx, t in enumerate(transactions, start=2):
        values = [
            t.created_at.strftime("%d-%m-%Y"),
            t.display_id,
            t.user.employee_id,
            t.user.name,
            float(t.amount),
            t.purpose,
            t.remarks or "",
            t.created_by,
        ]
        for c_idx, val in enumerate(values, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            if c_idx == 5:
                cell.number_format = CURRENCY_FORMAT

    last_col = get_column_letter(len(headers))
    ws.auto_filter.ref = f"A1:{last_col}{max(1, len(transactions) + 1)}"
    _autosize(ws, headers)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
